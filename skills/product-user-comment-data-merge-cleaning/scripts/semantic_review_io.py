"""Typed XLSX input and deterministic, staged AI-review artifacts. No AI calls."""
from __future__ import annotations

import csv
from contextlib import ExitStack
from datetime import date, datetime, time, timedelta
import json
import math
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

try:
    from .output_path_safety import ensure_output_paths_safe, atomic_output_path
except ImportError:
    from output_path_safety import ensure_output_paths_safe, atomic_output_path

HEADERS = ['评论日期', '评论内容', '产品名', '电商平台评分', '用户属性', '哈希ID',
           '点赞数', '子评论数/追评数', '一级评论', '二级评论', '三级评论']
AI_HEADERS = ['AI类别', '清洗处理', '保留原因']


def pack(value):
    if isinstance(value, (datetime, date, time)):
        return {'kind': type(value).__name__, 'value': value.isoformat()}
    if isinstance(value, timedelta):
        return {'kind': 'timedelta', 'value': value.total_seconds()}
    if value is None or type(value) in (str, bool, int, float):
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError('Non-finite numeric cell')
        return value
    raise ValueError('Unsupported cell type; do not silently convert it')


def unpack(value):
    if not isinstance(value, dict):
        return value
    kind, val = value['kind'], value['value']
    if kind == 'timedelta':
        return timedelta(seconds=val)
    return {'datetime': datetime, 'date': date, 'time': time}[kind].fromisoformat(val)


def snapshot(source: Path):
    if source.suffix.lower() != '.xlsx':
        raise ValueError('AI input must be the confirmed traditional final .xlsx, not CSV/raw exports')
    wb = load_workbook(source, read_only=True, data_only=False)
    rows, sheets = [], []
    try:
        for si, ws in enumerate(wb.worksheets):
            if list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True))) != HEADERS:
                raise ValueError(f'Invalid eleven-column schema in sheet index {si}')
            count = 0
            for rn, cells in enumerate(ws.iter_rows(min_row=2), 2):
                vals = [pack(c.value) for c in cells]
                key = f's{si}:r{rn}'
                h = vals[5]
                if h not in (None, '') and (type(h) is not str or re.fullmatch('[0-9a-f]{64}', h) is None):
                    raise ValueError(f'Invalid pseudonymous key at {key}')
                if vals[6] is None or isinstance(vals[6], str) and not vals[6].strip():
                    raise ValueError(f'Blank likes at {key}; run the traditional workflow first')
                rows.append({'id': key, 'sheet': ws.title, 'row': rn, 'values': vals,
                             'types': [c.data_type for c in cells],
                             'formats': [c.number_format for c in cells]})
                count += 1
            sheets.append({'name': ws.title, 'rows': count})
    finally:
        wb.close()
    return sheets, rows


def put(ws, rn, cn, value, data_type=None, number_format=None):
    if value == '' and data_type is None:
        value = None  # Excel serializes generated empty strings as blank cells.
    cell = ws.cell(rn, cn, unpack(value))
    # Strings originating from CSV or AI remain literal. Original real formulas
    # preserve type 'f' only when explicitly carried from the source snapshot.
    if data_type is not None:
        cell.data_type = data_type
    elif isinstance(cell.value, str):
        cell.data_type = 's'
    if number_format is not None:
        cell.number_format = number_format


def plain_row(ws, values):
    rn = ws.max_row + 1 if ws.cell(1, 1).value is not None else 1
    for cn, value in enumerate(values, 1):
        put(ws, rn, cn, value)


def source_row(ws, rn, row, offset=0):
    for i, value in enumerate(row['values']):
        put(ws, rn, i + 1 + offset, value, row['types'][i], row['formats'][i])


def style(wb):
    for ws in wb.worksheets:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='28556B')
            ws.column_dimensions[cell.column_letter].width = 20
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)


def make_workbooks(manifest, decisions, groups):
    final, review = Workbook(), Workbook()
    final.remove(final.active)
    for sheet in manifest['sheets']:
        plain_row(final.create_sheet(sheet['name']), HEADERS + AI_HEADERS)
    allws = review.active
    allws.title = '全量AI审查'
    aws = review.create_sheet('A类候选清单')
    review_headers = ['行键', '输入工作表', '输入行号'] + HEADERS + [
        '初始AI类别', '最终AI类别', 'K保护信号', '结构家族', 'R1链数', 'R2',
        '技术词命中', 'KOL词命中', 'R3成员', '标注理由', '全文复核处理', '复核理由', '结构证据']
    for ws in (allws, aws):
        plain_row(ws, review_headers)
    for row in manifest['rows']:
        d = decisions[row['id']]
        if d['final_label'] != 'A':
            ws = final[row['sheet']]
            rn = ws.max_row + 1
            source_row(ws, rn, row)
            for i, v in enumerate([d['final_label'], '保留', d['final_note']], 12):
                put(ws, rn, i, v)
        for ws in ([allws, aws] if d['label'] == 'A' else [allws]):
            rn = ws.max_row + 1
            for i, v in enumerate([row['id'], row['sheet'], row['row']], 1):
                put(ws, rn, i, v)
            source_row(ws, rn, row, 3)
            extra = [d['label'], d['final_label'], ','.join(k['code'] for k in d['k']),
                     d['family'], len(d['r1']), int(len(row['technical_hits']) >= 4),
                     ','.join(row['technical_hits']), ','.join(row['kol_hits']),
                     ','.join(d['r3']['members']) if d['r3'] else '', d['note'],
                     d.get('review_decision', ''), d.get('review_note', ''),
                     json.dumps({'evidence': d['evidence'], 'r1': d['r1'], 'r3': d['r3']}, ensure_ascii=False)]
            for i, v in enumerate(extra, 15):
                if isinstance(v, str) and len(v) > 32767:
                    raise ValueError('Review evidence exceeds XLSX cell limit; narrow evidence quotes')
                put(ws, rn, i, v)
    gws = review.create_sheet('R3模板群')
    plain_row(gws, ['组号', '成员行键', '不同非空哈希数', '证据'])
    for i, g in enumerate(groups, 1):
        plain_row(gws, [i, ','.join(g['members']), g['distinct_hashes'],
                       json.dumps(g['quotes'], ensure_ascii=False)])
    sws = review.create_sheet('清洗汇总')
    plain_row(sws, ['指标', '数值'])
    for k, v in [('输入行数', len(manifest['rows'])), ('批次数', len(manifest['batches'])),
                 ('初始A候选数', sum(d['label'] == 'A' for d in decisions.values())),
                 ('最终删除数', sum(d['final_label'] == 'A' for d in decisions.values())),
                 ('最终保留数', sum(d['final_label'] != 'A' for d in decisions.values())),
                 ('R3组数', len(groups)), ('运行摘要', manifest['run_id']),
                 ('说明', 'AI辅助判读；不证明真人、刷评或账号归属；CSV仅导出首个工作表')]:
        plain_row(sws, [k, v])
    style(final)
    style(review)
    return review, final


def text(value):
    if value is None:
        return ''
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def csv_rows(wb):
    return [[text(c.value) for c in row] for row in wb.worksheets[0].iter_rows()]


def cell_signature(cell):
    if cell.value is None:
        return None, 'empty', cell.number_format
    return pack(cell.value), cell.data_type, cell.number_format


def strictly_equal(expected, actual):
    if type(expected) is not type(actual):
        return False
    if isinstance(expected, dict):
        return expected.keys() == actual.keys() and all(
            strictly_equal(expected[key], actual[key])
            for key in expected
        )
    if isinstance(expected, (list, tuple)):
        return len(expected) == len(actual) and all(
            strictly_equal(left, right)
            for left, right in zip(expected, actual)
        )
    return expected == actual


def compare_workbook(path, expected):
    actual = load_workbook(path, read_only=True, data_only=False)
    try:
        if actual.sheetnames != expected.sheetnames:
            raise ValueError('Output worksheet order/name mismatch')
        for ws, es in zip(actual.worksheets, expected.worksheets):
            if (ws.max_row, ws.max_column) != (es.max_row, es.max_column):
                raise ValueError('Output dimensions mismatch')
            for row, exp in zip(ws.iter_rows(), es.iter_rows()):
                if not strictly_equal(
                    [cell_signature(c) for c in row],
                    [cell_signature(c) for c in exp],
                ):
                    raise ValueError('Output cell value/type/format mismatch')
    finally:
        actual.close()


def verify_outputs(manifest, decisions, groups, outputs):
    expected_review, expected_final = make_workbooks(manifest, decisions, groups)
    try:
        compare_workbook(outputs[0], expected_review)
        compare_workbook(outputs[1], expected_final)
        with Path(outputs[2]).open(encoding='utf-8-sig', newline='') as f:
            if list(csv.reader(f)) != csv_rows(expected_final):
                raise ValueError('CSV header/row/field mismatch')
    finally:
        expected_review.close()
        expected_final.close()


def export_outputs(manifest, decisions, groups, outputs, inputs, overwrite=False, confirmations=()):
    if len(outputs) != 3 or [p.suffix.lower() for p in outputs] != ['.xlsx', '.xlsx', '.csv']:
        raise ValueError('Require review XLSX, final XLSX and final CSV')
    if any(
        p.name.casefold().endswith(('.deletions.csv', '.summary.json'))
        for p in outputs
    ):
        raise ValueError('AI extension must not restore finalized traditional audit files')
    protected = [*inputs, *(p.with_suffix('.csv') for p in inputs if p.suffix.lower() == '.xlsx')]
    ensure_output_paths_safe(protected, outputs, overwrite=overwrite, overwrite_confirmations=confirmations)
    review, final = make_workbooks(manifest, decisions, groups)
    try:
        with ExitStack() as stack:
            staged = [stack.enter_context(atomic_output_path(p)) for p in outputs]
            review.save(staged[0])
            final.save(staged[1])
            with staged[2].open('w', encoding='utf-8-sig', newline='') as f:
                csv.writer(f).writerows(csv_rows(final))
            verify_outputs(manifest, decisions, groups, staged)
    finally:
        review.close()
        final.close()
