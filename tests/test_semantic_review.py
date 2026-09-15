from __future__ import annotations

import copy
import importlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from datetime import datetime, timedelta

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
HEADERS = ['评论日期', '评论内容', '产品名', '电商平台评分', '用户属性', '哈希ID',
           '点赞数', '子评论数/追评数', '一级评论', '二级评论', '三级评论']


class SemanticReviewTests(unittest.TestCase):
    def setUp(self):
        self.assertTrue((ROOT / 'tools/semantic_review.py').exists(), 'AI extension CLI is missing')
        self.contract = importlib.import_module('tools.semantic_review_contract')
        self.io = importlib.import_module('tools.semantic_review_io')
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name)
        self.src = self.dir / 'traditional.xlsx'

    def workbook(self, count=3, sheets=1):
        wb = Workbook()
        for s in range(sheets):
            ws = wb.active if s == 0 else wb.create_sheet()
            ws.title = f'表{s}'
            ws.append(HEADERS)
            for n in range(count):
                ws.append(['2026-09', '完整评论' * 60 + f'实际使用会打滑{n}', 'synthetic', 4.5,
                           None, format(n + 1, '064x') if n else None, 5,
                           None, '=1+1', None, None])
                ws.cell(n + 2, 9).data_type = 's'
        wb.save(self.src)
        wb.close()
        return self.src

    def manifest(self, count=3, sheets=1, terms=()):
        self.workbook(count, sheets)
        return self.contract.build_manifest(self.src, 'synthetic', 'bilibili', list(terms), [])

    def documents(self, manifest):
        return [{'run_id': manifest['run_id'], 'batch_id': b['id'], 'model': 'synthetic-test',
                 'prompt_sha256': manifest['prompt_sha256'],
                 'annotations': [{'id': key, 'label': 'C', 'k': [], 'family': '',
                                  'evidence': [], 'r1': [], 'r3': None, 'note': '合成测试保留'}
                                 for key in b['row_ids']]}
                for b in manifest['batches']]

    def quote(self, field='评论内容', quote='完整评论'):
        return {'field': field, 'quote': quote}

    def a_documents(self, manifest):
        docs = self.documents(manifest)
        docs[0]['annotations'][0].update(label='A', family='GSoc', evidence=[self.quote()])
        return docs

    def reviews(self, manifest, docs, decision='delete'):
        return {'run_id': manifest['run_id'], 'annotations_sha256': self.contract.digest(docs),
                'reviewer': 'independent-synthetic-reviewer', 'reviews': [
                    {'id': docs[0]['annotations'][0]['id'], 'decision': decision,
                     'full_text_reviewed': True, 'k': [], 'evidence': [self.quote()],
                     'note': '合成复核记录'}]}

    def test_row_boundaries_and_full_text(self):
        for count in (0, 1, 60, 61, 2000, 2310, 2400):
            with self.subTest(count=count):
                m = self.manifest(count)
                self.assertEqual(count, len(m['rows']))
                self.assertEqual(count, sum(len(b['row_ids']) for b in m['batches']))
                self.assertTrue(all(len(b['row_ids']) <= 60 for b in m['batches']))
                if count:
                    self.assertIn('实际使用会打滑', m['rows'][0]['values'][1])
                    self.assertIsNone(m['rows'][0]['values'][5])

    def test_source_binding_and_manifest_tamper(self):
        m = self.manifest()
        self.contract.check_manifest(self.src, m)
        bad = copy.deepcopy(m)
        bad['rows'][0]['values'][1] = 'changed'
        with self.assertRaises(ValueError):
            self.contract.check_manifest(self.src, bad)
        wb = load_workbook(self.src)
        wb.active['G2'] = 999
        wb.save(self.src)
        wb.close()
        with self.assertRaises(ValueError):
            self.contract.check_manifest(self.src, m)

    def test_small_batches_support_complete_long_text(self):
        self.workbook(5)
        m = self.contract.build_manifest(self.src, 'p', 'bilibili', [], [], batch_size=2)
        self.assertEqual([2, 2, 1], [len(b['row_ids']) for b in m['batches']])
        self.contract.check_manifest(self.src, m)
        for size in (0, 61, True):
            with self.assertRaises(ValueError):
                self.contract.build_manifest(self.src, 'p', 'bilibili', [], [], batch_size=size)

    def test_R1_positive_and_numeric_dates_formula_types(self):
        self.workbook(1)
        wb = load_workbook(self.src)
        ws = wb.active
        ws['B2'] = '功能甲效果甲功能乙效果乙'
        ws['A2'] = datetime(2026, 9, 14, 12, 30)
        ws['D2'] = '=2+2'
        ws['G2'] = '未识别的点赞文本'
        ws['J2'] = timedelta(hours=2)
        wb.save(self.src)
        wb.close()
        m = self.contract.build_manifest(self.src, 'p', 'bilibili', ['功能甲', '功能甲', '效果甲'], [])
        docs = self.documents(m)
        docs[0]['annotations'][0]['r1'] = [
            {'function': self.quote(quote='功能甲'), 'effect': self.quote(quote='效果甲')},
            {'function': self.quote(quote='功能乙'), 'effect': self.quote(quote='效果乙')}]
        d, g = self.contract.validate_annotations(m, docs, require_reviews=True)
        self.assertEqual(2, len(m['rows'][0]['technical_hits']))
        outputs = [self.dir / 'review.xlsx', self.dir / 'AI.xlsx', self.dir / 'AI.csv']
        self.io.export_outputs(m, d, g, outputs, [self.src])
        wb = load_workbook(outputs[1])
        self.assertEqual('f', wb.active['D2'].data_type)
        self.assertEqual('未识别的点赞文本', wb.active['G2'].value)
        wb.close()

    def test_empty_export(self):
        m = self.manifest(0)
        d, g = self.contract.validate_annotations(m, [], require_reviews=True)
        outputs = [self.dir / 'review.xlsx', self.dir / 'AI.xlsx', self.dir / 'AI.csv']
        self.io.export_outputs(m, d, g, outputs, [self.src])
        self.io.verify_outputs(m, d, g, outputs)

    def test_annotation_missing_duplicate_invalid_fields_and_empty_directory(self):
        m = self.manifest(61)
        docs = self.documents(m)
        self.contract.validate_annotations(m, docs)
        cases = [[], docs[:1], docs + docs[:1]]
        for key, value in [('label', 'D'), ('k', None), ('family', 'unknown'), ('note', 42),
                           ('id', True), ('r1', True)]:
            bad = copy.deepcopy(docs)
            bad[0]['annotations'][0][key] = value
            cases.append(bad)
        bad = copy.deepcopy(docs)
        del bad[0]['annotations'][0]['family']
        cases.append(bad)
        for bad in cases:
            with self.subTest(bad=str(bad)[:40]), self.assertRaises(ValueError):
                self.contract.validate_annotations(m, bad)

    def test_A_requires_full_review_and_no_K(self):
        m = self.manifest()
        docs = self.a_documents(m)
        with self.assertRaises(ValueError):
            self.contract.validate_annotations(m, docs, require_reviews=True)
        review = self.reviews(m, docs)
        decisions, _ = self.contract.validate_annotations(m, docs, review, require_reviews=True)
        self.assertEqual('A', decisions[m['rows'][0]['id']]['final_label'])
        review['reviews'][0]['decision'] = 'keep'
        decisions, _ = self.contract.validate_annotations(m, docs, review, require_reviews=True)
        self.assertEqual('B', decisions[m['rows'][0]['id']]['final_label'])
        review['reviews'][0]['decision'] = 'delete'
        review['reviews'][0]['k'] = [dict(code='K5', **self.quote(quote='会打滑'))]
        with self.assertRaises(ValueError):
            self.contract.validate_annotations(m, docs, review, require_reviews=True)
        docs[0]['annotations'][0]['k'] = [dict(code='K5', **self.quote(quote='会打滑'))]
        with self.assertRaises(ValueError):
            self.contract.validate_annotations(m, docs)

    def test_fake_evidence_and_stale_review_rejected(self):
        m = self.manifest()
        docs = self.a_documents(m)
        review = self.reviews(m, docs)
        docs[0]['annotations'][0]['note'] = '修改了标注'
        with self.assertRaises(ValueError):
            self.contract.validate_annotations(m, docs, review, require_reviews=True)
        docs[0]['annotations'][0]['evidence'] = [self.quote(quote='源表没有的文字')]
        with self.assertRaises(ValueError):
            self.contract.validate_annotations(m, docs)

    def test_R1_and_R3_are_actual_evidence(self):
        m = self.manifest(4)
        docs = self.documents(m)
        a = docs[0]['annotations'][1]
        a['r1'] = [{'function': self.quote(), 'effect': self.quote(quote='实际使用') }]
        with self.assertRaises(ValueError):
            self.contract.validate_annotations(m, docs)
        a['r1'] = []
        members = [r['id'] for r in m['rows'][1:]]
        a['r3'] = {'members': members, 'quotes': [dict(id=key, **self.quote()) for key in members]}
        _, groups = self.contract.validate_annotations(m, docs)
        self.assertEqual(3, groups[0]['distinct_hashes'])
        a['r3']['members'] = members[:2]
        a['r3']['quotes'] = a['r3']['quotes'][:2]
        with self.assertRaises(ValueError):
            self.contract.validate_annotations(m, docs)

    def test_schema_hash_and_source_types(self):
        self.workbook()
        for cell, value in [('A1', '错误表头'), ('F2', '123abc'), ('G2', None)]:
            self.workbook()
            wb = load_workbook(self.src)
            wb.active[cell] = value
            wb.save(self.src)
            wb.close()
            with self.assertRaises(ValueError):
                self.contract.build_manifest(self.src, 'p', 'bilibili', [], [])

    def test_export_verify_multisheet_typed_and_all_columns(self):
        m = self.manifest(3, 2)
        docs = self.a_documents(m)
        review = self.reviews(m, docs)
        decisions, groups = self.contract.validate_annotations(m, docs, review, require_reviews=True)
        outputs = [self.dir / 'review.xlsx', self.dir / 'AI.xlsx', self.dir / 'AI.csv']
        self.io.export_outputs(m, decisions, groups, outputs, [self.src])
        self.io.verify_outputs(m, decisions, groups, outputs)
        wb = load_workbook(outputs[1])
        self.assertEqual(['表0', '表1'], wb.sheetnames)
        self.assertEqual('s', wb['表0']['I2'].data_type)
        self.assertEqual('=1+1', wb['表0']['I2'].value)
        self.assertEqual(4.5, wb['表0']['D2'].value)
        wb.close()
        for col in range(1, 15):
            self.io.export_outputs(m, decisions, groups, outputs, [self.src], True, outputs)
            wb = load_workbook(outputs[1])
            wb['表0'].cell(2, col, 'CORRUPTED')
            wb.save(outputs[1])
            wb.close()
            with self.subTest(col=col), self.assertRaises(ValueError):
                self.io.verify_outputs(m, decisions, groups, outputs)

    def test_output_overlap_and_partial_overwrite_preflight(self):
        m = self.manifest()
        decisions, groups = self.contract.validate_annotations(m, self.documents(m), require_reviews=True)
        outputs = [self.dir / 'review.xlsx', self.dir / 'AI.xlsx', self.dir / 'AI.csv']
        outputs[2].write_text('untouched', encoding='utf-8')
        with self.assertRaises(ValueError):
            self.io.export_outputs(m, decisions, groups, outputs, [self.src], True, [])
        self.assertFalse(outputs[0].exists())
        self.assertEqual('untouched', outputs[2].read_text())
        with self.assertRaises(ValueError):
            self.io.export_outputs(m, decisions, groups, [outputs[0], self.src, outputs[2]], [self.src])

    def test_cannot_restore_traditional_logs_or_overwrite_companion_csv(self):
        m = self.manifest()
        d, g = self.contract.validate_annotations(m, self.documents(m), require_reviews=True)
        outputs = [self.dir / 'review.xlsx', self.dir / 'AI.xlsx', self.dir / 'old.deletions.csv']
        with self.assertRaises(ValueError):
            self.io.export_outputs(m, d, g, outputs, [self.src])
        self.src.with_suffix('.csv').write_text('traditional', encoding='utf-8')
        outputs[2] = self.src.with_suffix('.csv')
        with self.assertRaises(ValueError):
            self.io.export_outputs(m, d, g, outputs, [self.src], True, [outputs[2]])

    def test_cli_and_copied_skill(self):
        self.workbook(1)
        skill = self.dir / 'portable'
        shutil.copytree(ROOT / 'skills/product-user-comment-data-merge-cleaning', skill)
        cli = skill / 'scripts/semantic_review.py'
        manifest = self.dir / 'run.json'
        env = dict(os.environ, PYTHONUTF8='1')
        def run(*args):
            return subprocess.run([sys.executable, str(cli), *map(str, args)], cwd=self.dir,
                                  env=env, capture_output=True, text=True, encoding='utf-8')
        r = run('prepare', '--input', self.src, '--output', manifest, '--project', 'synthetic',
                '--platform', 'bilibili', '--confirm-traditional-output')
        self.assertEqual(0, r.returncode, r.stderr)
        m = json.loads(manifest.read_text(encoding='utf-8'))
        ann = self.dir / 'ann.json'
        ann.write_text(json.dumps(self.documents(m)[0], ensure_ascii=False), encoding='utf-8')
        common = ['--input', self.src, '--manifest', manifest, '--annotation', ann]
        outputs = ['--review-output', self.dir / 'review.xlsx', '--output', self.dir / 'AI.xlsx',
                   '--csv-output', self.dir / 'AI.csv']
        self.assertNotEqual(0, run('validate', '--input', self.src, '--manifest', manifest).returncode)
        r = run('export', *common, *outputs)
        self.assertEqual(0, r.returncode, r.stderr)
        r = run('verify', *common, *outputs)
        self.assertEqual(0, r.returncode, r.stderr)


if __name__ == '__main__':
    unittest.main()
